"""Genera el manual de usuario (HTML, PDF y DOCX) desde la evidencia Playwright."""

from __future__ import annotations

import html
import json
import subprocess
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image

from inspeccionar_aplicacion import chrome_path


ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
DATA_DIR = DOCS / "manual_usuario"
REPORT_JSON = DATA_DIR / "reporte_ejecucion.json"
AUDIT_JSON = DATA_DIR / "auditoria_formularios.json"
HTML_PATH = DOCS / "Manual_de_usuario_TransPort.html"
PDF_PATH = DOCS / "Manual_de_usuario_TransPort.pdf"
DOCX_PATH = DOCS / "Manual_de_usuario_TransPort.docx"
TEST_REPORT_PATH = DOCS / "Reporte_de_pruebas_manual.md"


USERS = [
    ("Administrador", "admin.demo", "TransPortAdmin2026!", "Configuración, auditoría y todos los módulos."),
    ("Gerente", "gerente.demo", "TransPortGerente2026!", "Gestión y supervisión operativa."),
    ("Empleado", "empleado.demo", "TransPortEmpleado2026!", "Operación interna asignada."),
    ("Operativo", "operativo.demo", "TransPortOperativo2026!", "Flujos operativos y logísticos."),
    ("Vendedor", "vendedor.demo", "TransPortVendedor2026!", "Catálogo, pedidos y comunicaciones."),
    ("Cliente", "cliente.demo", "TransPortCliente2026!", "Portal de una empresa cliente."),
    ("Cliente — una empresa", "cliente2.demo", "TransPortCliente22026!", "Alcance restringido a una empresa."),
    ("Cliente — varias empresas", "cliente3.demo", "TransPortCliente32026!", "Selección entre dos empresas."),
]

CHAPTER_DESCRIPTIONS = {
    "Sitio público": "Consulta de la portada y del catálogo sin iniciar sesión.",
    "Acceso y seguridad": "Autenticación, validación de credenciales y cierre de sesión.",
    "Perfiles y portales": "Menús iniciales según rol y alcance autorizado.",
    "Administración": "Usuarios, clientes, roles, giros, sedes y auditoría.",
    "Portal del cliente": "Pedidos, operaciones, reportes y comunicaciones de las empresas asociadas.",
    "Gestión operativa": "Páginas de entrada, listados, búsqueda y filtros de los módulos internos.",
    "Inventarios y bodegas": "Artículos, existencias y movimientos entre bodegas.",
    "Abastecimiento": "Solicitudes internas de abastecimiento y sus transiciones.",
    "Extracción minera": "Registro y seguimiento de jornadas de extracción.",
    "Orden de producción": "Creación de órdenes de transformación.",
    "Producción": "Seguimiento de estados y etapas de producción.",
    "Equipo de maquinaria": "Registro de maquinaria, contratos y mantenimiento.",
    "Transporte y GPS": "Vehículos, conductores, viajes, incidencias y telemetría.",
    "Transporte": "Actualización de viajes existentes.",
    "GPS": "Simulación de coordenadas y eventos externos.",
    "Proyecto de construcción": "Creación de obras, actividades, recursos y avances.",
    "Presupuesto": "Creación de presupuestos operativos.",
    "Finanzas": "Gastos, presupuestos e ingresos.",
    "Alerta": "Creación de avisos y recordatorios.",
    "Alertas": "Lectura, seguimiento y resolución de avisos.",
    "Comunicación interna": "Conversaciones internas y mensajes de seguimiento.",
    "Comunicacion": "Operaciones de conversación sobre registros creados.",
    "Integraciones y finanzas": "Simulación REST, Webhook y CSV, además de ingresos manuales.",
}


def load_data() -> tuple[dict, dict]:
    report = json.loads(REPORT_JSON.read_text(encoding="utf-8"))
    audit = json.loads(AUDIT_JSON.read_text(encoding="utf-8"))
    return report, audit


def manual_captures(report: dict) -> list[dict]:
    # Cada submit ya tiene una captura del resultado. La captura adicional que
    # safe() toma al registrar un fallo sería redundante dentro del manual.
    return [item for item in report["captures"] if not item["title"].startswith("Fallo -")]


def grouped_captures(captures: list[dict]) -> OrderedDict[str, list[dict]]:
    result: OrderedDict[str, list[dict]] = OrderedDict()
    for item in captures:
        result.setdefault(item["chapter"], []).append(item)
    return result


def failure_source_url(report: dict, failure: dict) -> str:
    expected = f"{failure['name']} - datos".casefold()
    for capture in reversed(report["captures"]):
        if capture["chapter"] == failure["chapter"] and capture["title"].casefold() == expected:
            return capture["url"]
    return failure.get("url") or "No disponible"


def generate_html(report: dict, audit: dict) -> None:
    captures = manual_captures(report)
    chapters = grouped_captures(captures)
    summary = report["summary"]
    failures = [item for item in report["checks"] if item["status"] == "fallo"]

    toc = "".join(
        f'<li><a href="#cap-{index}">{index}. {html.escape(chapter)}</a></li>'
        for index, chapter in enumerate(chapters, start=1)
    )
    credentials = "".join(
        "<tr>"
        f"<td>{html.escape(role)}</td><td><code>{html.escape(user)}</code></td>"
        f"<td><code>{html.escape(password)}</code></td><td>{html.escape(scope)}</td>"
        "</tr>"
        for role, user, password, scope in USERS
    )

    chapter_html = []
    for chapter_number, (chapter, items) in enumerate(chapters.items(), start=1):
        cards = []
        for step, item in enumerate(items, start=1):
            status_class = "failure" if item["outcome"] == "fallo" else "success"
            status_label = "Resultado con fallo" if item["outcome"] == "fallo" else "Evidencia"
            cards.append(
                f"""
                <article class="capture-card {status_class}">
                  <div class="capture-heading">
                    <div><span class="step">Paso {step}</span><h3>{html.escape(item['title'])}</h3></div>
                    <span class="status">{status_label}</span>
                  </div>
                  <p>{html.escape(item['description'])}</p>
                  <p class="route"><strong>Ruta observada:</strong> <code>{html.escape(item['url'])}</code></p>
                  <img src="manual_usuario/{html.escape(item['image'])}" alt="{html.escape(item['title'])}" loading="eager" />
                  <p class="caption">Figura {chapter_number}.{step}. {html.escape(item['title'])}</p>
                </article>
                """
            )
        chapter_html.append(
            f"""
            <section class="chapter" id="cap-{chapter_number}">
              <header class="chapter-title"><span>Capítulo {chapter_number}</span><h2>{html.escape(chapter)}</h2>
              <p>{html.escape(CHAPTER_DESCRIPTIONS.get(chapter, 'Procedimiento comprobado en la aplicación.'))}</p></header>
              {''.join(cards)}
            </section>
            """
        )

    failure_rows = "".join(
        f"<tr><td>{html.escape(item['chapter'])}</td><td>{html.escape(item['name'])}</td>"
        f"<td><code>{html.escape(failure_source_url(report, item))}</code></td><td>{html.escape(item['detail'])}</td></tr>"
        for item in failures
    )
    if failures or audit["summary"]["forms_without_antiforgery"]:
        validation_html = (
            '<div class="warning"><strong>Incidencias pendientes.</strong> '
            f'La auditoría encontró {audit["summary"]["forms_without_antiforgery"]} formularios POST '
            f'potencialmente afectados en {audit["summary"]["files_affected"]} vistas.</div>'
        )
        failures_html = (
            '<h3>Flujos que fallaron</h3><table><thead><tr><th>Módulo</th><th>Flujo</th>'
            f'<th>Ruta</th><th>Resultado</th></tr></thead><tbody>{failure_rows}</tbody></table>'
        )
    else:
        validation_html = (
            '<div class="success-note"><strong>Validación completa.</strong> '
            'Todos los flujos automatizados finalizaron correctamente y la auditoría no encontró '
            'formularios POST sin protección antifalsificación.</div>'
        )
        failures_html = '<h3>Flujos fallidos</h3><p>No se registraron fallos en esta ejecución.</p>'

    document = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8" />
<meta name="viewport" content="width=device-width, initial-scale=1" />
<title>Manual de usuario · TransPort</title>
<style>
  :root{{--ink:#102b3a;--muted:#56717f;--brand:#ef7d22;--line:#d9e3e8;--soft:#f3f7f8;--danger:#a92727;--ok:#26734d}}
  *{{box-sizing:border-box}} body{{margin:0;color:var(--ink);font:15px/1.5 Arial,sans-serif;background:#fff}}
  main{{max-width:1120px;margin:auto;padding:42px}} h1,h2,h3{{line-height:1.15;margin-top:0}} h1{{font-size:52px}} h2{{font-size:34px}} h3{{font-size:23px}}
  .cover{{min-height:92vh;display:flex;flex-direction:column;justify-content:center;border-bottom:12px solid var(--brand)}}
  .cover-mark{{font-size:24px;font-weight:800;color:var(--brand);letter-spacing:.08em}} .subtitle{{font-size:24px;color:var(--muted);max-width:720px}}
  .meta{{margin-top:70px;color:var(--muted)}} .intro,.toc,.credentials,.appendix{{padding:50px 0;border-bottom:1px solid var(--line)}}
  .toc ol{{columns:2;gap:40px}} .toc li{{margin:0 0 10px}} a{{color:#155d77}} table{{width:100%;border-collapse:collapse;margin:18px 0}}
  th,td{{padding:10px;border:1px solid var(--line);text-align:left;vertical-align:top}} th{{background:var(--soft)}} code{{font:13px Consolas,monospace;word-break:break-all}}
  .chapter{{padding-top:54px}} .chapter-title{{border-left:8px solid var(--brand);padding:6px 0 6px 22px;margin-bottom:40px}} .chapter-title>span,.step{{text-transform:uppercase;letter-spacing:.08em;font-weight:700;font-size:12px;color:var(--brand)}}
  .capture-card{{margin:0 0 54px;padding:24px;border:1px solid var(--line);border-radius:16px;break-inside:avoid;page-break-inside:avoid;background:white}}
  .capture-card.failure{{border-color:#e7a5a5}} .capture-heading{{display:flex;align-items:flex-start;justify-content:space-between;gap:20px}}
  .status{{font-size:12px;font-weight:700;border-radius:20px;padding:7px 11px;background:#e8f5ee;color:var(--ok)}} .failure .status{{background:#fdeaea;color:var(--danger)}}
  .route{{font-size:13px;color:var(--muted)}} img{{display:block;max-width:100%;max-height:860px;object-fit:contain;margin:18px auto;border:1px solid #cbd8dd;border-radius:8px}}
  .caption{{font-size:12px;text-align:center;color:var(--muted)}} .summary-grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:14px;margin:22px 0}}
  .metric{{background:var(--soft);padding:18px;border-radius:12px}} .metric strong{{display:block;font-size:30px}} .warning{{padding:18px;border-left:6px solid var(--danger);background:#fff1f1}}
  .success-note{{padding:18px;border-left:6px solid var(--ok);background:#e8f5ee}}
  @page{{size:A4;margin:14mm}} @media print{{main{{max-width:none;padding:0}} .cover{{min-height:260mm;page-break-after:always}} .intro,.toc,.credentials{{page-break-after:always}} .chapter{{page-break-before:always}} .capture-card{{page-break-before:always;margin:0;padding:10px;border:none}} .capture-card:first-of-type{{page-break-before:auto}} img{{max-height:190mm}} a{{color:inherit;text-decoration:none}}}}
  @media(max-width:760px){{main{{padding:22px}}h1{{font-size:38px}}.toc ol{{columns:1}}.summary-grid{{grid-template-columns:1fr 1fr}}}}
</style></head><body><main>
<section class="cover"><div class="cover-mark">TRANSPORT</div><h1>Manual de usuario</h1><p class="subtitle">Guía completa de navegación y operación del prototipo, elaborada a partir de pruebas reales con evidencia visual antes y después de enviar los formularios.</p><p class="meta">Versión documentada: {datetime.now():%d/%m/%Y}<br/>Entorno: http://localhost:8080<br/>Ejecución: {html.escape(report['run_id'])}</p></section>
<section class="intro"><h2>Cómo usar este manual</h2><p>Abra TransPort en <code>http://localhost:8080</code>. Cada capítulo muestra la pantalla, los datos introducidos y el resultado observado. Las capturas con borde rojo corresponden a una función que falló durante la prueba y se conservan como evidencia; consulte el apéndice de incidencias antes de demostrar esos flujos.</p>
<h3>Convenciones</h3><ul><li><strong>Evidencia:</strong> la pantalla se abrió o el formulario devolvió un resultado normal.</li><li><strong>Resultado con fallo:</strong> el navegador o la aplicación rechazó la operación.</li><li>Todos los datos creados por la prueba incluyen el prefijo <code>{html.escape(report['run_id'])}</code>.</li></ul></section>
<section class="toc"><h2>Contenido</h2><ol>{toc}<li><a href="#incidencias">Apéndice: resultado de pruebas e incidencias</a></li></ol></section>
<section class="credentials"><h2>Credenciales de demostración</h2><p>Estas cuentas pertenecen exclusivamente al entorno local de demostración.</p><table><thead><tr><th>Perfil</th><th>Usuario</th><th>Contraseña</th><th>Alcance</th></tr></thead><tbody>{credentials}</tbody></table></section>
{''.join(chapter_html)}
<section class="appendix" id="incidencias"><h2>Apéndice: resultado de las pruebas</h2><div class="summary-grid"><div class="metric"><strong>{summary['checks']}</strong>comprobaciones</div><div class="metric"><strong>{summary['passed']}</strong>correctas</div><div class="metric"><strong>{summary['failed']}</strong>fallidas</div><div class="metric"><strong>{summary['screenshots']}</strong>capturas</div></div>
{validation_html}
{failures_html}
<p>Consulte <code>Reporte_de_pruebas_manual.md</code>, <code>manual_usuario/reporte_ejecucion.json</code> y <code>manual_usuario/auditoria_formularios.json</code> para el detalle reproducible.</p></section>
</main></body></html>"""
    HTML_PATH.write_text(document, encoding="utf-8")


def add_docx_picture(document: Document, image_path: Path) -> None:
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    max_width = 6.4
    max_height = 7.4
    ratio = width_px / height_px
    width = min(max_width, max_height * ratio)
    height = width / ratio
    document.add_picture(str(image_path), width=Inches(width), height=Inches(height))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def generate_docx(report: dict, audit: dict) -> None:
    captures = manual_captures(report)
    chapters = grouped_captures(captures)
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    title = document.add_heading("Manual de usuario", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("TransPort · Prototipo integral")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Ejecución documentada: {report['run_id']}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    document.add_heading("Credenciales de demostración", 1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for index, heading in enumerate(("Perfil", "Usuario", "Contraseña", "Alcance")):
        table.rows[0].cells[index].text = heading
    for row in USERS:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    document.add_page_break()

    for chapter, items in chapters.items():
        document.add_heading(chapter, 1)
        document.add_paragraph(CHAPTER_DESCRIPTIONS.get(chapter, "Procedimiento comprobado en la aplicación."))
        for step, item in enumerate(items, start=1):
            document.add_heading(f"{step}. {item['title']}", 2)
            document.add_paragraph(item["description"])
            document.add_paragraph(f"Ruta: {item['url']}")
            add_docx_picture(document, DATA_DIR / item["image"])
            if item["outcome"] == "fallo":
                document.add_paragraph("Resultado observado: FALLO (consulte el apéndice técnico).")
            document.add_page_break()

    document.add_heading("Apéndice: resultado de pruebas", 1)
    summary = report["summary"]
    document.add_paragraph(
        f"Comprobaciones: {summary['checks']}. Correctas: {summary['passed']}. "
        f"Fallidas: {summary['failed']}. Capturas: {summary['screenshots']}."
    )
    if report["summary"]["failed"] == 0 and audit["summary"]["forms_without_antiforgery"] == 0:
        document.add_paragraph("Todos los flujos automatizados finalizaron correctamente y no se detectaron formularios POST sin token antifalsificación.")
    else:
        document.add_paragraph(
            f"Se detectaron {audit['summary']['forms_without_antiforgery']} formularios POST sin token "
            f"antifalsificación en {audit['summary']['files_affected']} vistas."
        )
    document.save(DOCX_PATH)


def generate_test_report(report: dict, audit: dict) -> None:
    failures = [item for item in report["checks"] if item["status"] == "fallo"]
    affected_files: OrderedDict[str, list[dict]] = OrderedDict()
    for finding in audit["findings"]:
        affected_files.setdefault(finding["file"], []).append(finding)
    lines = [
        "# Reporte de pruebas del manual de usuario",
        "",
        f"- Ejecución: `{report['run_id']}`",
        f"- Fecha: {report['generated_at']}",
        f"- URL: `{report['base_url']}`",
        f"- Comprobaciones: **{report['summary']['checks']}**",
        f"- Correctas: **{report['summary']['passed']}**",
        f"- Fallidas: **{report['summary']['failed']}**",
        f"- Capturas: **{report['summary']['screenshots']}**",
        "",
    ]
    if not failures and audit["summary"]["forms_without_antiforgery"] == 0:
        lines.extend([
            "## Resultado final",
            "",
            "Todos los flujos automatizados finalizaron correctamente. La auditoría estática tampoco encontró formularios POST sin token antifalsificación.",
            "",
        ])
    else:
        lines.extend([
            "## Flujos fallidos",
            "",
        ])
        for failure in failures:
            lines.extend(
                [
                    f"- **{failure['chapter']} — {failure['name']}**",
                    f"  - Ruta: `{failure_source_url(report, failure)}`",
                    f"  - Resultado: {failure['detail']}",
                ]
            )
        lines.extend(
            [
                "",
                "## Auditoría de formularios POST",
                "",
                f"Se encontraron **{audit['summary']['forms_without_antiforgery']} formularios** "
                f"potencialmente afectados en **{audit['summary']['files_affected']} vistas**.",
                "",
            ]
        )
    for file, findings in affected_files.items():
        lines.append(f"### `{file}`")
        lines.append("")
        for finding in findings:
            lines.append(f"- Línea {finding['line']}: `{finding['action']}`")
        lines.append("")
    lines.extend(
        [
            "## Evidencia reproducible",
            "",
            "- `manual_usuario/reporte_ejecucion.json`: cada comprobación y captura.",
            "- `manual_usuario/auditoria_formularios.json`: hallazgos estáticos.",
            "- `manual_usuario/capturas/`: imágenes antes y después de cada envío.",
            "- `../tools/manual_usuario/capturar_manual.py`: ejecución Playwright.",
        ]
    )
    TEST_REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def generate_pdf() -> None:
    command = [
        chrome_path(),
        "--headless=new",
        "--disable-gpu",
        "--no-sandbox",
        "--no-pdf-header-footer",
        "--run-all-compositor-stages-before-draw",
        f"--print-to-pdf={PDF_PATH}",
        HTML_PATH.resolve().as_uri(),
    ]
    subprocess.run(command, check=True, timeout=300)


def main() -> int:
    report, audit = load_data()
    generate_html(report, audit)
    generate_test_report(report, audit)
    generate_docx(report, audit)
    generate_pdf()
    for path in (HTML_PATH, PDF_PATH, DOCX_PATH, TEST_REPORT_PATH):
        print(f"Generado: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
